"""Read-only group document/opaque archive evidence; never execute attachments."""
import hashlib
import json
from pathlib import Path
import struct
import sys

from index_juku_assets import REPO, save_json, file_hash

DEPS = Path('C:/Users/theon/AppData/Local/Temp/juku-group-reader-20261019-deps')
sys.path.insert(0, str(DEPS))
import olefile
import py7zr
from docx import Document
from pypdf import PdfReader


def main():
    root = Path('D:/AM6B-Project/AM6B/资源')
    out = REPO/'docs/research/juku-full-20261019/assets/group-documents'
    records = []
    for source in sorted(root.iterdir()):
        ext = source.suffix.lower()
        if ext not in ('.docx', '.doc', '.pdf', '.temp'):
            continue
        sha = file_hash(source)
        row = dict(source=str(source), name=source.name, bytes=source.stat().st_size, sha256=sha)
        if ext == '.docx':
            document = Document(source)
            row.update(state='semantic-extraction-only',
                       layout='Not rendered: bundled LibreOffice unavailable; no layout conclusions',
                       paragraphs=[p.text for p in document.paragraphs],
                       tables=[[[c.text for c in r.cells] for r in t.rows] for t in document.tables])
        elif ext == '.pdf':
            document = PdfReader(source)
            row.update(state='page-text-extracted', pages=[dict(page=i+1,text=p.extract_text()) for i,p in enumerate(document.pages)])
        elif ext == '.doc':
            with olefile.OleFileIO(source) as document:
                streams = []
                for name in document.listdir():
                    data = document.openstream(name).read()
                    streams.append(dict(path='/'.join(name),bytes=len(data),sha256=hashlib.sha256(data).hexdigest()))
                row.update(state='legacy-OLE-stream-index-only', streams=streams,
                           limitation='Word binary text/layout not decoded; macros/embedded objects not executed')
        elif ext == '.temp':
            with source.open('rb') as handle:
                header = handle.read(32)
            if header[:6] != b'7z\xbc\xaf\x27\x1c':
                raise ValueError('Unexpected opaque archive type')
            _, offset, length, crc = struct.unpack('<IQQI', header[8:32])
            row.update(format='7z',next_header_offset=offset,next_header_bytes=length,
                       required_minimum_bytes=32+offset+length)
            try:
                with py7zr.SevenZipFile(source) as archive:
                    row.update(state='member-index-only', members=[dict(path=x.filename,bytes=x.uncompressed,
                                                                       directory=x.is_directory) for x in archive.list()])
            except py7zr.Bad7zFile as error:
                row.update(state='invalid-or-incomplete-archive', error=str(error),
                           limitation='Header points beyond actual file; no member payload recovered')
        target = out/(sha[:12]+'.json')
        save_json(target,row)
        records.append(dict(name=source.name,file=str(target.relative_to(out.parent)).replace('\\','/'),
                            sha256=sha,state=row['state'],pages=len(row.get('pages',[]))))
    save_json(out/'index.json',dict(reader='bundled Python; python-docx/pypdf; task-local olefile 0.47/py7zr 1.1.4',files=records))
    print(json.dumps(records,ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
